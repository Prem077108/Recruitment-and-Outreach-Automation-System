from __future__ import annotations

import shutil
from pathlib import Path

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from docx2pdf import convert
except ImportError:  # pragma: no cover
    convert = None

BASE_DIR = Path(__file__).resolve().parent
SCRATCH_DIR = BASE_DIR / "scratch"
SCRATCH_DIR.mkdir(exist_ok=True)


def extract_resume_text(path: str | Path) -> str:
    resume_path = Path(path)
    if resume_path.suffix.lower() == ".docx" and Document is not None:
        doc = Document(str(resume_path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    return resume_path.read_text(encoding="utf-8", errors="ignore") if resume_path.exists() else ""


def customize_resume(resume_path: str | Path, job_description: str, candidate_name: str, client_name: str = "demo") -> Path:
    source = Path(resume_path)
    if not source.exists():
        raise FileNotFoundError(f"Resume file not found: {source}")

    summary = (
        f"Candidate: {candidate_name}\n"
        f"Client: {client_name}\n"
        f"Job description focus: {job_description[:400]}\n\n"
        "This version emphasizes strong engineering depth, strong communication, and US C2C contract alignment."
    )

    if source.suffix.lower() == ".docx" and Document is not None:
        doc = Document(str(source))
        doc.add_paragraph(summary)
        target_docx = SCRATCH_DIR / f"{client_name}_{source.stem}_customized.docx"
        doc.save(target_docx)

        if convert is not None:
            target_pdf = target_docx.with_suffix(".pdf")
            try:
                convert(str(target_docx), str(target_pdf))
            except Exception:
                pass
            if target_pdf.exists():
                return target_pdf
        return target_docx

    target_pdf = SCRATCH_DIR / f"{client_name}_{source.stem}_customized.pdf"
    shutil.copy2(source, target_pdf)
    return target_pdf
